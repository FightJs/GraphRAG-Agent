"""文档解析服务 — 将上传的原始文件提取为纯文本，供 KG 抽取和问答使用"""
import re
from html.parser import HTMLParser
from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from bs4 import BeautifulSoup


def extract_text(file_path: str, file_format: str) -> tuple[str, int]:
    """返回 (纯文本, 页数/分片数估计)"""
    fmt = file_format.lower()
    if fmt == "pdf":
        return _extract_pdf(file_path)
    if fmt == "docx":
        return _extract_docx(file_path)
    if fmt == "xlsx":
        return _extract_xlsx(file_path)
    if fmt == "pptx":
        return _extract_pptx(file_path)
    if fmt == "html":
        return _extract_html(file_path)
    if fmt in ("txt", "md", "csv"):
        text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
        return text, max(1, text.count("\n") // 40 + 1)
    raise ValueError(f"5001:不支持解析的文件格式 {fmt}")


def _extract_pdf(file_path: str) -> tuple[str, int]:
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages), len(reader.pages)


def _extract_docx(file_path: str) -> tuple[str, int]:
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(paragraphs)
    return text, max(1, len(paragraphs) // 30 + 1)


def _extract_xlsx(file_path: str) -> tuple[str, int]:
    wb = load_workbook(file_path, data_only=True, read_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"# 表: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines), len(wb.worksheets)


def _extract_pptx(file_path: str) -> tuple[str, int]:
    prs = Presentation(file_path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"# 幻灯片 {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
    return "\n".join(lines), len(prs.slides)


def _extract_html(file_path: str) -> tuple[str, int]:
    html = Path(file_path).read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    text = soup.get_text("\n", strip=True)
    return text, max(1, text.count("\n") // 40 + 1)


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 150) -> list[str]:
    """按字符数滑动窗口分块，边界对齐完整占位符"""
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        end = _align_chunk_boundary(text, start, end, overlap)
        if end <= start:
            end = min(start + chunk_size, n)
        chunks.append(text[start:end])
        if end >= n:
            break
        start = max(0, end - overlap)
        # overlap 也不能切开占位符：若 start 落在占位符内则后移
        for m in _PLACEHOLDER_ANY_RE.finditer(text[max(0, start - 32) : min(n, start + 32)]):
            abs_start = max(0, start - 32) + m.start()
            abs_end = max(0, start - 32) + m.end()
            if abs_start < start < abs_end:
                start = abs_end
                break
    return chunks


# ── MinerU content_list 结构化解析 ─────────────────────────────────────────

_SKIP_BLOCK_TYPES = {"image", "equation", "interline_equation"}
_PLACEHOLDER_LINE_RE = re.compile(r"^\s*\{\{(TABLE|IMAGE):([a-z0-9_]+)\}\}\s*$")
_PLACEHOLDER_ANY_RE = re.compile(r"\{\{(TABLE|IMAGE):([a-z0-9_]+)\}\}")


def extract_media_refs(text: str) -> list[str]:
    """从 chunk 文本中提取占位符引用的 media_id 列表（去重保序）。"""
    if not text:
        return []
    seen: set[str] = set()
    refs: list[str] = []
    for _kind, mid in _PLACEHOLDER_ANY_RE.findall(text):
        if mid not in seen:
            seen.add(mid)
            refs.append(mid)
    return refs


def _align_chunk_boundary(text: str, start: int, end: int, overlap: int) -> int:
    """滑窗边界不得切开占位符；在 end 附近前向/后向微调到完整占位符之外。"""
    if end >= len(text):
        return end
    window = text[max(0, start - 64) : min(len(text), end + 64)]
    base = max(0, start - 64)
    for m in _PLACEHOLDER_ANY_RE.finditer(window):
        abs_start = base + m.start()
        abs_end = base + m.end()
        if abs_start < end < abs_end:
            # 切在占位符内部：优先把 end 推到占位符后；若过长则拉回占位符前
            if abs_end - start <= 2000 and abs_end <= len(text):
                return min(abs_end + 1, len(text))
            return max(abs_start, start + 1)
    return end


def as_text(value) -> str:
    """MinerU 字段可能是 str / list[str] / None；统一转成去空白字符串。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (list, tuple)):
        parts = [str(item).strip() for item in value if item is not None and str(item).strip()]
        return " ".join(parts)
    return str(value).strip()


def _block_text(block: dict) -> str:
    # 官方规范字段是 content；兼容 text / table_body / table_html，以及 list 形态
    for key in ("content", "text", "table_body", "table_html"):
        text = as_text(block.get(key))
        if text:
            return text
    return ""


def _html_table_to_md(html_content: str, caption: str = "") -> str:
    if not html_content or "<table" not in html_content.lower():
        # 已是 Markdown 表则原样返回
        if html_content.lstrip().startswith("|"):
            return (f"{caption}\n" if caption else "") + html_content.strip() + "\n"
        return ""

    class TableParser(HTMLParser):
        def __init__(self):
            super().__init__()
            self.rows: list[list[str]] = []
            self.current_row: list[str] = []
            self.in_table = False
            self.in_row = False
            self.in_cell = False
            self.current_cell: list[str] = []

        def handle_starttag(self, tag, attrs):
            if tag == "table":
                self.in_table = True
            elif tag == "tr" and self.in_table:
                self.in_row = True
                self.current_row = []
            elif tag in ("th", "td") and self.in_row:
                self.in_cell = True
                self.current_cell = []

        def handle_endtag(self, tag):
            if tag == "table":
                self.in_table = False
            elif tag == "tr" and self.in_row:
                self.in_row = False
                if self.current_row:
                    self.rows.append(self.current_row)
            elif tag in ("th", "td") and self.in_cell:
                self.in_cell = False
                self.current_row.append("".join(self.current_cell).strip())

        def handle_data(self, data):
            if self.in_cell:
                self.current_cell.append(data)

    parser = TableParser()
    try:
        parser.feed(html_content)
    except Exception:
        return ""
    if not parser.rows:
        return ""
    lines = []
    if caption:
        lines.append(caption)
    for i, row in enumerate(parser.rows):
        lines.append("| " + " | ".join(row) + " |")
        if i == 0:
            lines.append("| " + " | ".join(["---"] * len(row)) + " |")
    return "\n".join(lines) + "\n"


def _render_block(
    block: dict,
    include_tables: bool = True,
    media_table_id: str | None = None,
    media_image_id: str | None = None,
    media_mode: str = "legacy",
) -> str:
    block_type = block.get("type", "")
    content = _block_text(block)
    if block_type == "image":
        if media_mode == "enhanced" and media_image_id:
            caption = as_text(block.get("img_caption")) or as_text(block.get("image_caption"))
            parts = []
            if caption:
                parts.append(caption)
            parts.append(f"{{{{IMAGE:{media_image_id}}}}}")
            return "\n\n".join(parts) + "\n"
        return ""
    if block_type in _SKIP_BLOCK_TYPES:
        return ""
    if block_type == "text":
        level = block.get("text_level") or 0
        if level > 0:
            return f"{'#' * level} {content}\n"
        return f"{content}\n" if content else ""
    if block_type == "table":
        caption = as_text(block.get("table_caption"))
        if media_mode == "enhanced" and media_table_id:
            parts = []
            if caption:
                parts.append(caption)
            parts.append(f"{{{{TABLE:{media_table_id}}}}}")
            footer = as_text(block.get("table_footnote"))
            if footer:
                parts.append(footer)
            return "\n\n".join(parts) + "\n"
        if include_tables:
            md = _html_table_to_md(content, caption)
            footer = as_text(block.get("table_footnote"))
            if md and footer:
                md = md.rstrip("\n") + f"\n{footer}\n"
            return md
        return ""
    if block_type == "list":
        return f"{content}\n" if content else ""
    if block_type == "code":
        return f"```\n{content}\n```\n" if content else ""
    return f"{content}\n" if content else ""


def content_list_to_markdown(
    blocks: list[dict],
    include_tables: bool = True,
    table_ids: dict[int, str] | None = None,
    image_ids: dict[int, str] | None = None,
) -> str:
    """将 MinerU content_list 转为可读 Markdown 全文。

    table_ids / image_ids：content_list 下标 → media_id；提供时走占位符渲染。
    """
    table_ids = table_ids or {}
    image_ids = image_ids or {}
    media_mode = "enhanced" if (table_ids or image_ids) else "legacy"
    parts = []
    for i, block in enumerate(blocks or []):
        if not isinstance(block, dict):
            continue
        rendered = _render_block(
            block,
            include_tables=include_tables,
            media_table_id=table_ids.get(i),
            media_image_id=image_ids.get(i),
            media_mode=media_mode if (table_ids.get(i) or image_ids.get(i) or block.get("type") not in ("table", "image")) else "legacy",
        )
        # 未命中增强映射的 table/image：table 走 include_tables，image 跳过
        if block.get("type") == "table" and i not in table_ids:
            rendered = _render_block(block, include_tables=include_tables, media_mode="legacy")
        if block.get("type") == "image" and i not in image_ids:
            rendered = ""
        if rendered.strip():
            parts.append(rendered)
    return "\n".join(parts).strip()


def content_list_to_chunks(
    blocks: list[dict],
    split_by: str = "page",
    chunk_size: int = 1200,
    overlap: int = 150,
    include_tables: bool = True,
    table_ids: dict[int, str] | None = None,
    image_ids: dict[int, str] | None = None,
) -> list[dict]:
    """按 MinerU 结构块生成带页码/章节元数据的分块。

    返回 [{"text", "page_idx", "section", "chunk_index", "media_refs", "content_type", "media_mode"}]
    split_by: page（默认，按页聚合后必要时再切）| section（H1 章节）| full（整篇再滑窗）
    """
    if not blocks:
        return []
    split_by = (split_by or "page").lower()
    if split_by not in ("page", "section", "full"):
        split_by = "page"
    table_ids = table_ids or {}
    image_ids = image_ids or {}
    media_mode = "enhanced" if (table_ids or image_ids) else "legacy"

    def _render(i: int, block: dict) -> str:
        if block.get("type") == "table" and i not in table_ids:
            return _render_block(block, include_tables=include_tables, media_mode="legacy")
        if block.get("type") == "image" and i not in image_ids:
            return ""
        return _render_block(
            block,
            include_tables=include_tables,
            media_table_id=table_ids.get(i),
            media_image_id=image_ids.get(i),
            media_mode=media_mode if (table_ids.get(i) or image_ids.get(i)) else "legacy",
        )

    if split_by == "full":
        parts = []
        for i, block in enumerate(blocks):
            if not isinstance(block, dict):
                continue
            rendered = _render(i, block)
            if rendered.strip():
                parts.append(rendered)
        md = "\n".join(parts).strip()
        texts = chunk_text(md, chunk_size=chunk_size, overlap=overlap) if md else []
        return [
            {
                "text": t,
                "page_idx": None,
                "section": None,
                "chunk_index": i,
                "media_refs": extract_media_refs(t),
                "content_type": "text",
                "media_mode": media_mode,
            }
            for i, t in enumerate(texts)
        ]

    if split_by == "page":
        pages: dict[int, list[tuple[int, dict]]] = {}
        for i, block in enumerate(blocks):
            if not isinstance(block, dict):
                continue
            if block.get("type", "") in _SKIP_BLOCK_TYPES and block.get("type") != "image":
                continue
            if block.get("type") == "image" and i not in image_ids:
                continue
            page_idx = block.get("page_idx")
            if not isinstance(page_idx, int):
                page_idx = 0
            pages.setdefault(page_idx, []).append((i, block))
        raw_units: list[dict] = []
        for page_idx in sorted(pages):
            texts = [_render(i, b) for i, b in pages[page_idx]]
            raw_units.append(
                {
                    "text": "\n".join(t for t in texts if t.strip()).strip(),
                    "page_idx": page_idx,
                    "section": None,
                }
            )
    else:  # section
        raw_units = []
        current: list[str] = []
        current_title = ""
        current_pages: set[int] = set()
        for i, block in enumerate(blocks):
            if not isinstance(block, dict):
                continue
            btype = block.get("type", "")
            if btype in _SKIP_BLOCK_TYPES and btype != "image":
                continue
            if btype == "image" and i not in image_ids:
                continue
            page_idx = block.get("page_idx") if isinstance(block.get("page_idx"), int) else None
            if btype == "text" and (block.get("text_level") or 0) == 1:
                if current and current_title:
                    raw_units.append(
                        {
                            "text": "\n".join(current).strip(),
                            "page_idx": min(current_pages) if current_pages else None,
                            "section": current_title,
                        }
                    )
                current_title = _block_text(block)[:30]
                current = [f"# {_block_text(block)}\n"]
                current_pages = {page_idx} if page_idx is not None else set()
            else:
                rendered = _render(i, block)
                if rendered.strip():
                    current.append(rendered)
                if page_idx is not None:
                    current_pages.add(page_idx)
        if current and current_title:
            raw_units.append(
                {
                    "text": "\n".join(current).strip(),
                    "page_idx": min(current_pages) if current_pages else None,
                    "section": current_title,
                }
            )

    chunks: list[dict] = []
    for unit in raw_units:
        text = (unit["text"] or "").strip()
        if not text:
            continue
        pieces = (
            chunk_text(text, chunk_size=chunk_size, overlap=overlap)
            if len(text) > chunk_size
            else [text]
        )
        for piece in pieces:
            chunks.append(
                {
                    "text": piece,
                    "page_idx": unit["page_idx"],
                    "section": unit["section"],
                    "chunk_index": len(chunks),
                    "media_refs": extract_media_refs(piece),
                    "content_type": "text",
                    "media_mode": media_mode,
                }
            )
    return chunks


def plain_text_to_chunks(
    text: str,
    chunk_size: int = 1200,
    overlap: int = 150,
) -> list[dict]:
    """纯文本路径（pypdf 等）统一切成与结构化分块相同的字典格式"""
    return [
        {
            "text": t,
            "page_idx": None,
            "section": None,
            "chunk_index": i,
            "media_refs": [],
            "content_type": "text",
            "media_mode": "legacy",
        }
        for i, t in enumerate(chunk_text(text, chunk_size=chunk_size, overlap=overlap))
    ]


def normalize_chunks(raw: list) -> list[dict]:
    """兼容旧版 list[str] 与新版 list[dict] 分块文件；补 media_refs 等默认值"""
    result: list[dict] = []
    for i, item in enumerate(raw or []):
        if isinstance(item, str):
            result.append(
                {
                    "text": item,
                    "page_idx": None,
                    "section": None,
                    "chunk_index": i,
                    "media_refs": [],
                    "content_type": "text",
                    "media_mode": "legacy",
                }
            )
        elif isinstance(item, dict):
            text = item.get("text") or item.get("content") or ""
            result.append(
                {
                    "text": text,
                    "page_idx": item.get("page_idx"),
                    "section": item.get("section"),
                    "chunk_index": item.get("chunk_index", i),
                    "media_refs": item.get("media_refs") or extract_media_refs(text),
                    "content_type": item.get("content_type") or "text",
                    "media_mode": item.get("media_mode") or ("enhanced" if item.get("media_refs") else "legacy"),
                }
            )
    return result


def chunks_as_texts(chunks: list[dict] | list[str]) -> list[str]:
    return [c["text"] for c in normalize_chunks(chunks)]
